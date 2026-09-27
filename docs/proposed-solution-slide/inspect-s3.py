import io
import sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pptx import Presentation
prs = Presentation('slides/output/SIH26146_PPT_combined.pptx')
print('Slide count:', len(prs.slides))
slide = prs.slides[2]
for i, shape in enumerate(slide.shapes):
    txt = ''
    if shape.has_text_frame:
        txt = shape.text_frame.text[:100].replace('\n', ' | ')
    print(f'{i:3d} type={shape.shape_type} name={shape.name!r} text={txt!r}')