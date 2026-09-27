"""Inspect slide 2 of SIH26146_PPT_combined.pptx to find the architecture image state."""
import io
import os
import sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = os.path.dirname(os.path.abspath(__file__))
DECK_PATH = os.path.join(HERE, 'slides', 'output', 'SIH26146_PPT_combined.pptx')

prs = Presentation(DECK_PATH)
slide2 = prs.slides[1]

print('Total shapes:', len(slide2.shapes))
print('Slide width x height:', prs.slide_width / 914400.0, 'x', prs.slide_height / 914400.0)
print()

for i, shape in enumerate(slide2.shapes):
    try:
        l = shape.left / 914400.0 if shape.left is not None else None
        t = shape.top / 914400.0 if shape.top is not None else None
        w = shape.width / 914400.0 if shape.width is not None else None
        h = shape.height / 914400.0 if shape.height is not None else None
    except Exception as e:
        l = t = w = h = 'err'
    st = int(shape.shape_type) if shape.shape_type is not None else 'n/a'
    name = shape.name
    has_text = False
    text_preview = ''
    if shape.has_text_frame:
        has_text = True
        try:
            text_preview = shape.text_frame.text[:60].replace('\n', ' | ').encode('ascii', 'replace').decode('ascii')
        except Exception:
            text_preview = '<unreadable>'
    print(f'  [{i:3d}] type={st:>2} name="{name}" x={l} y={t} w={w} h={h} text={text_preview!r}')

# Check for image relationships on the slide
print()
print('=== Slide 2 relationships ===')
for rel_id, rel in slide2.part.rels.items():
    print(f'  {rel_id}: {rel.reltype.split("/")[-1]} -> {rel.target_ref}')