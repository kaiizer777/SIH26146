"""
fix-slide-02-v2.py
==================

Surgical fix for the three remaining issues on slide 2 (Proposed Solution)
of SIH26146_PPT_combined.pptx, identified after the previous worker's
enhancement pass:

  1. Title "Proposed Solution" was bumped to 38pt (and textbox widened to
     6.0in). Restore to original 24pt and original width 5.166in so it
     matches the visual scale of slides 3, 4, 5.

  2. The "ARCHITECTURE  --  BITCOIN AML INTELLIGENCE PIPELINE" pill banner
     overlaps the architecture image's own header (the pill sits on top of
     the image at y=1.22 and hides the image's "Bitcoin AML Intelligence
     Pipeline" title). Remove the pill entirely -- the image's own title
     is sufficient.

  3. Pillar textboxes are too short for their content -- PILLAR 1 text
     "...Neo4j 5.26 + GDS 2.13" is truncated, PILLAR 2 "...on CPU alone."
     is cut, PILLAR 4 ends mid-content. Resize all 5 pillar textboxes
     (plus their card backdrops and color stripes) to a uniform h=1.08in
     with touching tops, fitting in y in [1.40, 6.80] (just above the
     footer bar at y=6.951).

The script is idempotent: it does NOT remove any unrelated Enhance_*
shapes (background dots, footer chip, corner accents, title tint band,
title accent rule, etc.). It only touches the specific shapes that need
to change:

  - Text 4 (title)
  - Enhance_S2_ArchCaptionPill + Enhance_S2_ArchCaptionText
  - Text 9..Text 13 (pillar textboxes) + Enhance_S2_PillarCard_* +
    Enhance_S2_PillarStripe_*

Constraints honored:
  - Canvas 13.333 x 7.5 in.
  - ASCII-only Python string literals.
  - Slides 1, 3, 4, 5 untouched.

Run:
    python fix-slide-02-v2.py
"""

import io
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pptx import Presentation
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DECK_PATH = os.path.join(HERE, 'slides', 'output', 'SIH26146_PPT_combined.pptx')

# --- Target geometry ---
TITLE_LEFT = 3.324
TITLE_TOP = 0.554
TITLE_WIDTH = 5.166        # original; previous worker widened to 6.000
TITLE_HEIGHT = 0.700       # unchanged
TITLE_FONT_PT = 24         # original; previous worker bumped to 38

# Pillars: all h=1.08, touching tops, span [1.40, 6.80].
PILLAR_LEFT = 0.180
PILLAR_WIDTH = 5.500
PILLAR_STRIPE_WIDTH = 0.100
PILLAR_HEIGHT = 1.08
PILLAR_FIRST_TOP = 1.40
PILLAR_STEP = 1.08         # touching (gap = 0)

# Names used by the prior enhancement pass.
PILLAR_NAMES = (
    ('Text 9',  'THE SYSTEM'),
    ('Text 10', 'PILLAR 1'),
    ('Text 11', 'PILLAR 2'),
    ('Text 12', 'PILLAR 3'),
    ('Text 13', 'PILLAR 4'),
)
PILLAR_CARD_NAMES = ('Enhance_S2_PillarCard_1', 'Enhance_S2_PillarCard_2',
                     'Enhance_S2_PillarCard_3', 'Enhance_S2_PillarCard_4',
                     'Enhance_S2_PillarCard_5')
PILLAR_STRIPE_NAMES = ('Enhance_S2_PillarStripe_1', 'Enhance_S2_PillarStripe_2',
                       'Enhance_S2_PillarStripe_3', 'Enhance_S2_PillarStripe_4',
                       'Enhance_S2_PillarStripe_5')


def find_by_name(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def set_title_size(title_shape, size_pt, width_in):
    title_shape.width = Inches(width_in)
    tf = title_shape.text_frame
    for para in tf.paragraphs:
        for run in para.runs:
            run.font.size = Pt(size_pt)
            # Drop any letter-spacing applied by the prior worker so the
            # title reads at its natural visual weight.
            from lxml import etree
            NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
            rPr = run._r.find('{%s}rPr' % NS_A)
            if rPr is not None and 'spc' in rPr.attrib:
                del rPr.attrib['spc']


def main():
    print('Opening:', DECK_PATH)
    prs = Presentation(DECK_PATH)
    slide2 = prs.slides[1]

    # -- 1. Title fix ---------------------------------------------------
    title = find_by_name(slide2, 'Text 4')
    if title is None:
        raise RuntimeError('Title "Text 4" not found')
    assert abs(title.left / 914400.0 - TITLE_LEFT) < 0.05
    assert abs(title.top / 914400.0 - TITLE_TOP) < 0.05
    set_title_size(title, TITLE_FONT_PT, TITLE_WIDTH)
    print('Title -> %dpt, width=%.3fin' % (TITLE_FONT_PT, TITLE_WIDTH))

    # -- 2. Remove the overlapping architecture pill --------------------
    pill = find_by_name(slide2, 'Enhance_S2_ArchCaptionPill')
    pill_text = find_by_name(slide2, 'Enhance_S2_ArchCaptionText')
    if pill is not None:
        remove_shape(pill)
        print('Removed Enhance_S2_ArchCaptionPill')
    if pill_text is not None:
        remove_shape(pill_text)
        print('Removed Enhance_S2_ArchCaptionText')

    # -- 3. Resize pillar textboxes + their card backdrops + stripes ----
    for i, ((tb_name, _), card_name, stripe_name) in enumerate(
        zip(PILLAR_NAMES, PILLAR_CARD_NAMES, PILLAR_STRIPE_NAMES)
    ):
        top_in = PILLAR_FIRST_TOP + i * PILLAR_STEP

        # 3a. Pillar textbox.
        tb = find_by_name(slide2, tb_name)
        if tb is None:
            raise RuntimeError('Pillar textbox %s not found' % tb_name)
        tb.left = Inches(PILLAR_LEFT)
        tb.top = Inches(top_in)
        tb.width = Inches(PILLAR_WIDTH)
        tb.height = Inches(PILLAR_HEIGHT)

        # 3b. Card backdrop (rounded rectangle behind the pillar).
        card = find_by_name(slide2, card_name)
        if card is not None:
            card.left = Inches(PILLAR_LEFT)
            card.top = Inches(top_in)
            card.width = Inches(PILLAR_WIDTH)
            card.height = Inches(PILLAR_HEIGHT)

        # 3c. Color stripe on the left of the pillar.
        stripe = find_by_name(slide2, stripe_name)
        if stripe is not None:
            stripe.left = Inches(PILLAR_LEFT)
            stripe.top = Inches(top_in)
            stripe.width = Inches(PILLAR_STRIPE_WIDTH)
            stripe.height = Inches(PILLAR_HEIGHT)

        bottom_in = top_in + PILLAR_HEIGHT
        print('Pillar %d -> top=%.3f h=%.3f bottom=%.3f' %
              (i + 1, top_in, PILLAR_HEIGHT, bottom_in))

    last_bottom = PILLAR_FIRST_TOP + len(PILLAR_NAMES) * PILLAR_STEP
    print('Last pillar bottom y=%.3f (footer bar at 6.951)' % last_bottom)
    assert last_bottom <= 6.95, 'Pillars overflow footer'

    prs.save(DECK_PATH)
    print('Wrote:', DECK_PATH)


if __name__ == '__main__':
    main()
