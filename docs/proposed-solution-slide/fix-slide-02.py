"""Focused fix for slide 2 (Proposed Solution).

Repairs:
  1. Architecture image hidden: Shape 14 (panel) sits AFTER Image 1 in spTree,
     so the panel renders on top of the image. Move panel to immediately
     BEFORE the image so the image is visible.
  2. Title "Proposed Solution" wraps to 2 lines (40pt @ width 5.166") and
     overlaps the ARCHITECTURE pill (which sits at y=0.85-1.13).
     - Set title font size to 38pt and width to 6.000" so it fits on one line.
     - Remove the misplaced Enhance_S2_ArchCaptionPill/Text.
     - Re-add the pill as a banner overlay at the top of the architecture image
       panel (y=1.22), spanning the panel width.

Idempotent and limited to slide 2.
"""

import io
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DECK_PATH = os.path.join(HERE, 'slides', 'output', 'SIH26146_PPT_combined.pptx')

NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

INK = '2563EB'
WHITE = 'FFFFFF'


def emu_in(v):
    return int(round(v * 914400))


def find_by_name(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


def find_by_text_starts(slide, prefix):
    for shape in slide.shapes:
        try:
            if shape.has_text_frame and shape.text_frame.text.strip().startswith(prefix):
                return shape
        except Exception:
            pass
    return None


def remove_all_named(slide, name):
    n = 0
    for shape in list(slide.shapes):
        if shape.name == name:
            shape._element.getparent().remove(shape._element)
            n += 1
    return n


def set_run_font_size(text_frame, size_pt):
    for para in text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(size_pt)


def set_run_letter_spacing(text_frame, hundredths_pt):
    for para in text_frame.paragraphs:
        for run in para.runs:
            rPr = run._r.find('{%s}rPr' % NS_A)
            if rPr is None:
                rPr = etree.SubElement(run._r, '{%s}rPr' % NS_A)
                run._r.insert(0, rPr)
            rPr.set('spc', str(int(hundredths_pt)))


def main():
    print('Opening:', DECK_PATH)
    prs = Presentation(DECK_PATH)
    slide = prs.slides[1]

    # --- 1. Locate key shapes by name / content (stable across re-runs). ---
    title = find_by_text_starts(slide, 'Proposed Solution')
    # The title text is exactly "Proposed Solution" (no leading whitespace),
    # but the pillar text also contains "Proposed Solution -- SIH26146".
    # Use exact-match check.
    if title is not None:
        if title.text_frame.text.strip() != 'Proposed Solution':
            title = None
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip() == 'Proposed Solution':
                    title = shape
                    break
    panel = None
    arch_image = None
    for shape in slide.shapes:
        if shape.name == 'Shape 14':
            panel = shape
        elif shape.name == 'Image 1':
            arch_image = shape
    print('Title:', title.name if title else None)
    print('Panel:', panel.name if panel else None)
    print('Arch image:', arch_image.name if arch_image else None)
    if panel is None or arch_image is None or title is None:
        raise RuntimeError('Missing required shape')

    # --- 2. Fix the title so it fits on one line. ---
    title.width = Inches(6.000)
    title.height = Inches(0.700)
    set_run_font_size(title.text_frame, 38)
    set_run_letter_spacing(title.text_frame, 100)
    print('Title resized to 6.000 x 0.700 in, font 38pt.')

    # --- 3. Remove the misplaced ARCHITECTURE pill (and its text frame). ---
    n1 = remove_all_named(slide, 'Enhance_S2_ArchCaptionPill')
    n2 = remove_all_named(slide, 'Enhance_S2_ArchCaptionText')
    print('Removed Enhance_S2_ArchCaptionPill:', n1, ', Text:', n2)

    # --- 4. Fix panel/image z-order: panel must be BEFORE image in spTree. ---
    panel_elem = panel._element
    img_elem = arch_image._element
    spTree = panel_elem.getparent()
    siblings = list(spTree)
    panel_idx = siblings.index(panel_elem)
    img_idx = siblings.index(img_elem)
    print('Before fix: panel_idx=%d img_idx=%d' % (panel_idx, img_idx))
    if panel_idx > img_idx:
        spTree.remove(panel_elem)
        img_elem.addprevious(panel_elem)
        # verify
        siblings = list(spTree)
        panel_idx = siblings.index(panel_elem)
        img_idx = siblings.index(img_elem)
        print('After fix:  panel_idx=%d img_idx=%d' % (panel_idx, img_idx))
    else:
        print('Z-order already correct.')

    # --- 5. Re-add the pill as a banner overlay at the top of the panel. ---
    # Banner position: spans the panel width, sits at the very top inside the
    # panel so it overlays the architecture image's own title bar area.
    pill_x = 5.85
    pill_y = 1.22
    pill_w = 7.30
    pill_h = 0.30
    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(pill_x), Inches(pill_y),
        Inches(pill_w), Inches(pill_h),
    )
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor.from_string(INK)
    pill.line.color.rgb = RGBColor.from_string(INK)
    pill.line.width = Pt(0.5)
    # adjust geometry: roundRect with adj=50000 (50%)
    prst = pill._element.find('.//{%s}prstGeom' % NS_A)
    if prst is not None:
        prst.set('prst', 'roundRect')
        for av in list(prst.findall('{%s}avLst' % NS_A)):
            prst.remove(av)
        av_lst = etree.SubElement(prst, '{%s}avLst' % NS_A)
        gd = etree.SubElement(av_lst, '{%s}gd' % NS_A)
        gd.set('name', 'adj')
        gd.set('fmla', 'val 50000')
    try:
        pill.name = 'Enhance_S2_ArchCaptionPill'
    except Exception:
        pass

    # Text frame for the pill
    pill_text = slide.shapes.add_textbox(
        Inches(pill_x), Inches(pill_y),
        Inches(pill_w), Inches(pill_h),
    )
    tf = pill_text.text_frame
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = 'ARCHITECTURE  --  BITCOIN AML INTELLIGENCE PIPELINE'
    r.font.name = 'Arial'
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(WHITE)
    set_run_letter_spacing(tf, 200)
    try:
        pill_text.name = 'Enhance_S2_ArchCaptionText'
    except Exception:
        pass

    # Move both new shapes to be AFTER the arch image in spTree (so they sit
    # on top of the image as a banner overlay). The panel must remain BEFORE
    # the image, so the order in spTree is:
    #   ... panel ... image ... pill ... pill_text ... corners ...
    # Pill should be after image so it overlays the image content.
    img_elem = arch_image._element
    pill_elem = pill._element
    pt_elem = pill_text._element
    # Insert pill right after the image, then pill_text right after the pill.
    img_elem.addnext(pill_elem)
    pill_elem.addnext(pt_elem)

    print('Pill re-added at y=%.2f (banner overlay at top of image panel).' % pill_y)

    # --- 6. Save. ---
    prs.save(DECK_PATH)
    print('Saved:', DECK_PATH)


if __name__ == '__main__':
    main()