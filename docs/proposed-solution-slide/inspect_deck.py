"""Inspect the combined deck structure (encoding-safe)."""
import io
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from pptx import Presentation  # noqa: E402

OUT_DIR = os.path.join(
    r"C:\Users\bari2\Desktop\SIH26146\docs\proposed-solution-slide\slides",
    "output",
)
COMBINED_PATH = os.path.join(OUT_DIR, "SIH26146_PPT_combined.pptx")

prs = Presentation(COMBINED_PATH)
print("Total slides:", len(prs.slides))
print("Slide dims inches:", prs.slide_width / 914400, "x", prs.slide_height / 914400)
for idx in range(len(prs.slides)):
    s = prs.slides[idx]
    print("--- Slide", idx + 1, "(" + str(len(s.shapes)) + " shapes)")
    for sh_i, shape in enumerate(s.shapes):
        kind = str(shape.shape_type)
        text = ""
        if shape.has_text_frame:
            t = shape.text_frame.text.strip().splitlines()
            if t:
                text = t[0][:60]
        ln = None
        try:
            ln = shape.line
        except Exception:
            ln = None
        ln_w = None
        ln_color = None
        if ln is not None:
            try:
                ln_w = ln.width
            except Exception:
                pass
            try:
                ln_color = ln.color.rgb if ln.color and ln.color.type else "n/a"
            except Exception:
                ln_color = "?"
        ln_w_pt = ln_w / 12700 if ln_w is not None else "no-line"
        print(
            f"  [{sh_i}] {kind} id={shape.shape_id} name={shape.name!r} "
            f"pos=({shape.left/914400:.3f}, {shape.top/914400:.3f}, "
            f"{shape.width/914400:.3f}x{shape.height/914400:.3f}) "
            f"line_w_pt={ln_w_pt} line_color={ln_color} text={text!r}"
        )
