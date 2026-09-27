import io
import sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pptx import Presentation
from pptx.util import Pt
prs = Presentation('slides/output/SIH26146_PPT_combined.pptx')
slide = prs.slides[2]
for i, shape in enumerate(slide.shapes):
    txt = ''
    if shape.has_text_frame:
        txt = shape.text_frame.text[:100].replace('\n', ' | ')
    sizes = []
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    sizes.append(str(r.font.size.pt))
    print(f'{i:3d} name={shape.name!r} sizes={sizes} text={txt!r}')