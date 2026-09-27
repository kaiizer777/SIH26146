"""Debug script to trace what happens in enhance_cards."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from lxml import etree

NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'


def add_shape_after(slide, target_shape, shape_kind,
                    left_in, top_in, width_in, height_in, name=None):
    new_shape = slide.shapes.add_shape(
        shape_kind,
        Inches(left_in), Inches(top_in),
        Inches(width_in), Inches(height_in),
    )
    new_elem = new_shape._element
    target_elem = target_shape._element
    target_elem.addnext(new_elem)
    if name:
        try:
            new_shape.name = name
        except Exception:
            pass
    return new_shape


def dump_near(slide, target, before_after):
    spTree = target._element.getparent()
    children = list(spTree)
    target_idx = children.index(target._element)
    print(f'\n{before_after}: target={target.name} at index {target_idx}, total {len(children)}')
    for i in range(max(0, target_idx-2), min(len(children), target_idx+10)):
        elem = children[i]
        name_elem = elem.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        name = name_elem.get('name', 'NONAME') if name_elem is not None else elem.tag.split('}')[1]
        print(f'  [{i}] {name}')


prs = Presentation('slides/output/SIH26146_PPT_combined_backup_pre_enhance2.pptx')
slide = prs.slides[3]

shape9 = None
for s in slide.shapes:
    if s.name == 'Shape 9':
        shape9 = s
        break

dump_near(slide, shape9, 'Initial')

# Mimic enhance_cards for card 0
# add_card_icon: IconBg, IconGearOuter, IconGearInner
s = add_shape_after(slide, shape9, MSO_SHAPE.OVAL, 0.38, 1.16, 0.30, 0.30, name='TestIconBg')
dump_near(slide, shape9, 'After IconBg')

s = add_shape_after(slide, shape9, MSO_SHAPE.OVAL, 0.43, 1.21, 0.20, 0.20, name='TestIconGearOuter')
dump_near(slide, shape9, 'After IconGearOuter')

s = add_shape_after(slide, shape9, MSO_SHAPE.OVAL, 0.47, 1.25, 0.12, 0.12, name='TestIconGearInner')
dump_near(slide, shape9, 'After IconGearInner')
