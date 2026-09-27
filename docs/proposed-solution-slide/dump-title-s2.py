"""Dump the title XML to understand its state."""
import io
import os
import sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
DECK_PATH = os.path.join(HERE, 'slides', 'output', 'SIH26146_PPT_combined.pptx')

prs = Presentation(DECK_PATH)
slide2 = prs.slides[1]

# Find the title text frame (Text 4 with "Proposed Solution")
for shape in slide2.shapes:
    if shape.has_text_frame and 'Proposed Solution' in shape.text_frame.text and 'PILLAR' not in shape.text_frame.text and '--' not in shape.text_frame.text and 'SIH26146' not in shape.text_frame.text:
        print('Found title:')
        print('Name:', shape.name)
        print('Position:', shape.left/914400, shape.top/914400, shape.width/914400, shape.height/914400)
        print('Text:', repr(shape.text_frame.text))
        print()
        xml = etree.tostring(shape._element, pretty_print=True).decode('utf-8')
        print(xml)
        break

# Find the pill
print('=' * 60)
for shape in slide2.shapes:
    if shape.name == 'Enhance_S2_ArchCaptionPill':
        print('Pill shape:')
        xml = etree.tostring(shape._element, pretty_print=True).decode('utf-8')
        print(xml)

# Find the pill text
print('=' * 60)
for shape in slide2.shapes:
    if shape.name == 'Enhance_S2_ArchCaptionText':
        print('Pill text:')
        xml = etree.tostring(shape._element, pretty_print=True).decode('utf-8')
        print(xml)

# Find the panel
print('=' * 60)
for shape in slide2.shapes:
    if shape.name == 'Shape 14':
        print('Panel (Shape 14):')
        xml = etree.tostring(shape._element, pretty_print=True).decode('utf-8')
        print(xml[:3000])

# Find the architecture image
print('=' * 60)
for shape in slide2.shapes:
    if shape.name == 'Image 1':
        print('Architecture image (Image 1):')
        xml = etree.tostring(shape._element, pretty_print=True).decode('utf-8')
        print(xml)