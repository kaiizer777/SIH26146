"""Check current spTree z-order around the architecture image."""
import io
import os
import sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
DECK_PATH = os.path.join(HERE, 'slides', 'output', 'SIH26146_PPT_combined.pptx')

prs = Presentation(DECK_PATH)
slide2 = prs.slides[1]

# Get spTree children in order
spTree = slide2.shapes._spTree
print('spTree children order:')
for i, child in enumerate(list(spTree)):
    nvProps = child.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
    if nvProps is not None:
        name = nvProps.get('name', '?')
        id_ = nvProps.get('id', '?')
        tag = child.tag.split('}')[-1]
        print(f'  [{i:3d}] <{tag}> id={id_} name="{name}"')

# Check relationships and image targets
print()
print('Slide 2 relationships:')
for rel_id, rel in slide2.part.rels.items():
    target = rel.target_ref
    rtype = rel.reltype.split('/')[-1]
    print(f'  {rel_id}: {rtype} -> {target}')
    if 'image' in rtype:
        try:
            img_part = rel.target_part
            print(f'    blob size: {len(img_part.blob)} bytes')
            print(f'    content type: {img_part.content_type}')
        except Exception as e:
            print(f'    err: {e}')

# Check if image part referenced by Image 1 (rId4) actually has data
print()
print('Verify rId4 -> image part:')
try:
    img_part = slide2.part.related_parts['rId4']
    print(f'  blob size: {len(img_part.blob)}')
except Exception as e:
    print(f'  err: {e}')

# List media files in the package
print()
print('Package media:')
import zipfile
with zipfile.ZipFile(DECK_PATH, 'r') as z:
    for n in z.namelist():
        if 'media' in n or 'image' in n.lower():
            info = z.getinfo(n)
            print(f'  {n}: {info.file_size} bytes')