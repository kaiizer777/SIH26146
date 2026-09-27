"""Verify line spacing on slide 4 card body bullets."""
import io
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from pptx import Presentation  # noqa: E402

OUT_DIR = os.path.join(
    r"C:\Users\bari2\Desktop\SIH26146\docs\proposed-solution-slide\slides",
    "output",
)
DECK_PATH = os.path.join(OUT_DIR, "SIH26146_PPT_combined.pptx")

prs = Presentation(DECK_PATH)
slide4 = prs.slides[3]

target_card_indices = [10, 17, 23, 30, 36, 42]
target_card_rects = []
for idx in target_card_indices:
    shp = slide4.shapes[idx]
    target_card_rects.append((
        shp.left, shp.top, shp.left + shp.width, shp.top + shp.height,
    ))

for shape in slide4.shapes:
    if not shape.has_text_frame:
        continue
    try:
        t_left = shape.left
        t_top = shape.top
        t_right = t_left + shape.width
        t_bottom = t_top + shape.height
    except Exception:
        continue
    is_body = False
    is_title = False
    for (l, t, r, b) in target_card_rects:
        if not (t_left >= l - 1000 and t_right <= r + 1000):
            continue
        title_bar_h_emu = int(0.50 * 914400)
        if t_top < t + title_bar_h_emu:
            is_title = True
        elif t_top >= t and t_bottom <= b:
            is_body = True
    if is_body:
        line_spacings = []
        for p in shape.text_frame.paragraphs:
            line_spacings.append(p.line_spacing)
        text_preview = shape.text_frame.text[:60].replace("\n", " | ")
        print(f"BODY  [{shape.shape_id}] {text_preview!r} line_spacing={line_spacings}")
    elif is_title:
        line_spacings = [p.line_spacing for p in shape.text_frame.paragraphs]
        text_preview = shape.text_frame.text[:60].replace("\n", " | ")
        print(f"TITLE [{shape.shape_id}] {text_preview!r} line_spacing={line_spacings}")
