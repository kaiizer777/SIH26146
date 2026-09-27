"""
Combine two PPTX decks:
  - Result slide 1 = v4 deck's slide 1 (the new enhanced cover)
  - Result slides 2..N+1 = v3 deck's slides 1..N (the original rich content)

Strategy:
  1. Open v3 (3 slides), open v4 (4 slides).
  2. Append v4's slide 1 to v3 (becomes slide 4 in v3's package).
  3. For each picture shape in the new slide, copy the image blob from v4's
     package into v3's package via Package.get_or_add_image_part(), create a
     new relationship, and rewrite the blip's rId to point at the new
     relationship.
  4. Reorder sldIdLst so the new cover slide is at position 0.
  5. Save as SIH26146_PPT_combined.pptx.
"""

import io
import os
import shutil
from copy import deepcopy

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(HERE, 'slides', 'output')
V3_PATH = os.path.join(OUT_DIR, 'SIH26146_PPT_v3.pptx')
V4_PATH = os.path.join(OUT_DIR, 'SIH26146_PPT_v4.pptx')
OUT_PATH = os.path.join(OUT_DIR, 'SIH26146_PPT_combined.pptx')

NS_A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
NS_R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def main():
    # Step 1: copy v3 -> output, then open both
    shutil.copy(V3_PATH, OUT_PATH)
    combined = Presentation(OUT_PATH)
    src = Presentation(V4_PATH)

    src_slide = src.slides[0]

    # Step 2: append src_slide to combined
    new_slide = combined.slides.add_slide(combined.slide_layouts[0])

    # Strip default placeholders
    for ph in list(new_slide.placeholders):
        ph._element.getparent().remove(ph._element)

    # Deep-copy every shape element from src_slide into new_slide
    for shape in src_slide.shapes:
        new_elem = deepcopy(shape._element)
        new_slide.shapes._spTree.append(new_elem)

    # Step 3: copy image parts for any picture shapes
    src_image_parts = {}
    for shape in src_slide.shapes:
        if shape.shape_type == 13:  # PICTURE
            blip = shape._element.find(f'.//{NS_A}blip')
            if blip is None:
                continue
            src_rId = blip.get(f'{NS_R}embed')
            src_image_parts[src_rId] = shape.image

    src_to_new_rid = {}
    for src_rId, image in src_image_parts.items():
        blob = image.blob
        # Register the image blob in combined's package.
        # Package.get_or_add_image_part accepts a file-like object.
        img_stream = io.BytesIO(blob)
        image_part = combined.part.package.get_or_add_image_part(img_stream)

        # Relate the new slide to the (possibly new) image part.
        new_rId = new_slide.part.relate_to(image_part, RT.IMAGE)
        src_to_new_rid[src_rId] = new_rId

    # Rewrite blip rIds in the new slide to point at the new relationships
    for shape in new_slide.shapes:
        if shape.shape_type == 13:
            blip = shape._element.find(f'.//{NS_A}blip')
            if blip is None:
                continue
            old_rId = blip.get(f'{NS_R}embed')
            if old_rId in src_to_new_rid:
                blip.set(f'{NS_R}embed', src_to_new_rid[old_rId])

    # Step 4: reorder sldIdLst so new_slide is at position 0
    sldIdLst = combined.slides._sldIdLst
    sldIds = list(sldIdLst)
    new_sld_id_elem = sldIds[-1]

    sldIdLst.remove(new_sld_id_elem)
    sldIdLst.insert(0, new_sld_id_elem)

    # Step 5: save
    combined.save(OUT_PATH)
    print(f'Wrote: {OUT_PATH}')
    print(f'Total slides: {len(combined.slides)}')
    for i, s in enumerate(combined.slides, 1):
        first_text = ''
        for shape in s.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                first_text = shape.text_frame.text.strip().splitlines()[0][:70]
                break
        print(f'  Slide {i}: {first_text}')


if __name__ == '__main__':
    main()
