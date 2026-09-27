"""Inspect card 1 text styling and a few other card texts for comparison."""
import sys
from pptx import Presentation
from lxml import etree

DECK = r"C:\Users\bari2\Desktop\SIH26146\docs\proposed-solution-slide\slides\output\SIH26146_PPT_combined.pptx"

prs = Presentation(DECK)
slide1 = prs.slides[0]

# Card 1 = Text 11, Card 2 = Text 15, Card 3 = Text 19, etc.
text_shapes_to_check = [
    ("Text 11", "Card 1 - Problem Statement ID"),
    ("Text 15", "Card 2 - Title"),
    ("Text 19", "Card 3 - Theme"),
    ("Text 23", "Card 4 - Category"),
    ("Text 27", "Card 5 - Team ID"),
    ("Text 31", "Card 6 - Team Name"),
]

NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

def describe_runs(shape):
    print("Shape name:", shape.name)
    print("  Text content:", repr(shape.text_frame.text))
    for pi, para in enumerate(shape.text_frame.paragraphs):
        print("  Para %d alignment=%s runs=%d" % (pi, para.alignment, len(para.runs)))
        for ri, run in enumerate(para.runs):
            print("    Run %d: text=%r" % (ri, run.text))
            print("      font.name=%s size=%s bold=%s color=%s" % (
                run.font.name, run.font.size, run.font.bold,
                run.font.color.rgb if run.font.color and run.font.color.type else None,
            ))
            rPr = run._r.find('{%s}rPr' % NS_A)
            if rPr is not None:
                xml_str = etree.tostring(rPr, pretty_print=True).decode()
                for line in xml_str.split("\n"):
                    print("      rPr:", line)

for shape_name, label in text_shapes_to_check:
    print("\n=== %s ===" % label)
    found = None
    for shape in slide1.shapes:
        if shape.name == shape_name:
            found = shape
            break
    if found is None:
        print("  NOT FOUND")
        continue
    describe_runs(found)
