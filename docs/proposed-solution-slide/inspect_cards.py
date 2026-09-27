"""Inspect fills and existing effects on slides 1 & 4."""
import io
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

OUT_DIR = os.path.join(
    r"C:\Users\bari2\Desktop\SIH26146\docs\proposed-solution-slide\slides",
    "output",
)
COMBINED_PATH = os.path.join(OUT_DIR, "SIH26146_PPT_combined.pptx")

prs = Presentation(COMBINED_PATH)

# Get slide 4 specifically
slide4 = prs.slides[3]
print("=== Slide 4 fills and effect info ===")
for idx, shape in enumerate(slide4.shapes):
    fill_info = "n/a"
    try:
        fill = shape.fill
        fill_type = str(fill.type)
        fill_info = f"type={fill_type}"
        try:
            fill_info += f" fore_color={fill.fore_color.rgb}"
        except Exception:
            pass
    except Exception as e:
        fill_info = f"err:{e}"
    # Look for existing effectLst
    sp_pr = shape._element.find(qn("p:spPr"))
    has_effect = "no"
    if sp_pr is not None:
        effect_lst = sp_pr.find(qn("a:effectLst"))
        if effect_lst is not None:
            has_effect = "yes"
            for child in effect_lst:
                tag = child.tag.split("}")[1]
                attrs = {k.split("}")[-1]: v for k, v in child.attrib.items()}
                print(f"  [{idx}] effect child: <a:{tag}> attrs={attrs} kids={len(child)}")
                for kc in child:
                    kt = kc.tag.split("}")[1]
                    print(f"     - <a:{kt}>")
    text = ""
    if shape.has_text_frame:
        t = shape.text_frame.text.strip().splitlines()
        if t:
            text = t[0][:50]
    print(f"  [{idx}] name={shape.name!r} fill={fill_info} has_effectLst={has_effect} text={text!r}")

# Now slide 1 specifically for fills and effects
slide1 = prs.slides[0]
print()
print("=== Slide 1 fills ===")
for idx, shape in enumerate(slide1.shapes):
    fill_info = "n/a"
    try:
        fill = shape.fill
        fill_type = str(fill.type)
        fill_info = f"type={fill_type}"
        try:
            fill_info += f" fore_color={fill.fore_color.rgb}"
        except Exception:
            pass
    except Exception as e:
        fill_info = f"err:{e}"
    sp_pr = shape._element.find(qn("p:spPr"))
    has_effect = "no"
    if sp_pr is not None:
        effect_lst = sp_pr.find(qn("a:effectLst"))
        if effect_lst is not None:
            has_effect = "yes"
    text = ""
    if shape.has_text_frame:
        t = shape.text_frame.text.strip().splitlines()
        if t:
            text = t[0][:40]
    print(
        f"  [{idx}] name={shape.name!r} "
        f"fill={fill_info} has_effectLst={has_effect} text={text!r}"
    )
