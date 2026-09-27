"""Inspect slide 1 of the combined deck to understand current state."""
import os
import sys
from pptx import Presentation
from pptx.util import Emu

DECK = r"C:\Users\bari2\Desktop\SIH26146\docs\proposed-solution-slide\slides\output\SIH26146_PPT_combined.pptx"

prs = Presentation(DECK)
print("Total slides:", len(prs.slides))
slide1 = prs.slides[0]
print("\n=== Slide 1 shapes (count=%d) ===" % len(slide1.shapes))

for i, shape in enumerate(slide1.shapes):
    name = shape.name
    try:
        left = Emu(shape.left).inches if shape.left is not None else None
        top = Emu(shape.top).inches if shape.top is not None else None
        width = Emu(shape.width).inches if shape.width is not None else None
        height = Emu(shape.height).inches if shape.height is not None else None
        pos = "(%.3f, %.3f) %.3f x %.3f" % (left, top, width, height)
    except Exception:
        pos = "(pos unknown)"
    text = ""
    if shape.has_text_frame:
        text = repr(shape.text_frame.text[:80])
    print("  [%3d] %-40s %-30s %s" % (i, name, pos, text))
