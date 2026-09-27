from pptx import Presentation
from lxml import etree

prs = Presentation('slides/output/SIH26146_PPT_combined_backup_pre_enhance2.pptx')
slide = prs.slides[3]
spTree = slide.shapes._spTree
for i, child in enumerate(spTree):
    tag = child.tag.split('}')[1]
    if tag == 'extLst':
        print(f'extLst at index {i}')
print('Last child tag:', spTree[-1].tag)
print('Last 5 children tags:')
for i in range(max(0, len(spTree)-5), len(spTree)):
    name_elem = spTree[i].find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
    name = name_elem.get('name', 'NONAME') if name_elem is not None else spTree[i].tag.split('}')[1]
    tag = spTree[i].tag.split('}')[1]
    print(f'  [{i}] {name} (tag={tag})')
